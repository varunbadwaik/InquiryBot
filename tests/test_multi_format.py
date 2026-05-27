import tempfile
import unittest
import csv
from pathlib import Path
from database import load_pdf_paths

class MultiFormatTests(unittest.TestCase):
    """Verify loading and parsing for TXT, MD, CSV, and DOCX formats."""

    def test_load_txt_file(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            txt_path = Path(directory) / "test_doc.txt"
            txt_path.write_text("Hello, this is a plain text document.", encoding="utf-8")
            
            docs = load_pdf_paths([txt_path])
            self.assertEqual(len(docs), 1)
            self.assertEqual(docs[0].page_content, "Hello, this is a plain text document.")
            self.assertEqual(docs[0].metadata["document_name"], "test_doc.txt")
            self.assertEqual(docs[0].metadata["page_number"], 1)

    def test_load_md_file(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            md_path = Path(directory) / "test_doc.md"
            md_path.write_text("# Markdown Title\nThis is a md document.", encoding="utf-8")
            
            docs = load_pdf_paths([md_path])
            self.assertEqual(len(docs), 1)
            self.assertEqual(docs[0].page_content, "# Markdown Title\nThis is a md document.")
            self.assertEqual(docs[0].metadata["document_name"], "test_doc.md")

    def test_load_csv_file(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            csv_path = Path(directory) / "test_doc.csv"
            with open(csv_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(["Name", "Price", "Category"])
                writer.writerow(["Item A", "100", "Widget"])
                writer.writerow(["Item B", "200", "Gadget"])
            
            docs = load_pdf_paths([csv_path])
            self.assertEqual(len(docs), 1)
            self.assertIn("Name: Item A, Price: 100, Category: Widget", docs[0].page_content)
            self.assertIn("Name: Item B, Price: 200, Category: Gadget", docs[0].page_content)
            self.assertEqual(docs[0].metadata["document_name"], "test_doc.csv")

    def test_load_docx_file(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            docx_path = Path(directory) / "test_doc.docx"
            import docx
            doc = docx.Document()
            doc.add_paragraph("This is a Word paragraph.")
            # Add table
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Table Cell Content"
            doc.save(str(docx_path))
            
            docs = load_pdf_paths([docx_path])
            self.assertEqual(len(docs), 1)
            self.assertIn("This is a Word paragraph.", docs[0].page_content)
            self.assertIn("Table Cell Content", docs[0].page_content)
            self.assertEqual(docs[0].metadata["document_name"], "test_doc.docx")

    def test_load_zip_file(self) -> None:
        import zipfile
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            zip_path = Path(directory) / "test_archive.zip"
            inner_txt = "Content inside zip file."
            inner_csv = "ColA,ColB\nValA,ValB"
            
            with zipfile.ZipFile(zip_path, "w") as z:
                z.writestr("inner_folder/file1.txt", inner_txt)
                z.writestr("file2.csv", inner_csv)
                
            docs = load_pdf_paths([zip_path])
            self.assertEqual(len(docs), 2)
            
            # Check txt document inside ZIP
            txt_doc = next(d for d in docs if d.metadata["document_name"].endswith("file1.txt"))
            self.assertEqual(txt_doc.page_content, "Content inside zip file.")
            self.assertEqual(txt_doc.metadata["parent_document"], "test_archive.zip")
            self.assertEqual(txt_doc.metadata["document_name"], "test_archive.zip / inner_folder/file1.txt")
            
            # Check csv document inside ZIP
            csv_doc = next(d for d in docs if d.metadata["document_name"].endswith("file2.csv"))
            self.assertIn("ColA: ValA, ColB: ValB", csv_doc.page_content)
            self.assertEqual(csv_doc.metadata["parent_document"], "test_archive.zip")
            self.assertEqual(csv_doc.metadata["document_name"], "test_archive.zip / file2.csv")

    def test_zip_slip_prevention(self) -> None:
        import zipfile
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            zip_path = Path(directory) / "malicious.zip"
            with zipfile.ZipFile(zip_path, "w") as z:
                z.writestr("../escaped.txt", "escaping")
            
            with self.assertRaises(ValueError) as context:
                load_pdf_paths([zip_path])
            self.assertIn("Malicious ZIP entry detected", str(context.exception))


if __name__ == "__main__":
    unittest.main()
